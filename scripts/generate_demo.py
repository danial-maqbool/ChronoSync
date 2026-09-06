"""Generate synthetic input fixtures under ignored data/demo; never use personal files."""
import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from backend.demo import examples
from docx import Document
import pymupdf

target=Path('data/demo');target.mkdir(parents=True,exist_ok=True)
for name,text,_ in examples(): (target/name).write_text(text,encoding='utf-8')
doc=Document();doc.add_heading('Synthetic university course outline',0)
for line in examples()[0][1].splitlines(): doc.add_paragraph(line)
doc.save(target/'course-outline.docx')
pdf=pymupdf.open();page=pdf.new_page();page.insert_textbox((60,60,535,700),'SYNTHETIC COURSE OUTLINE\n\n'+examples()[0][1],fontsize=13);pdf.save(target/'course-outline.pdf')
(target/'interview.eml').write_text('From: recruiter@example.test\nTo: applicant@example.test\nDate: Thu, 3 Sep 2026 10:00:00 +0500\nSubject: Synthetic interview\nContent-Type: text/plain; charset=utf-8\n\nInterview tomorrow at 3 PM.',encoding='utf-8')
print('Synthetic fixtures generated in data/demo')
