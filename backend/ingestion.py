"""Local parsers. Segment evidence is never rewritten by event normalization."""

import csv
import io
import json
import re
from email import policy
from email.parser import BytesParser
from email.utils import parsedate_to_datetime
from pathlib import Path
from bs4 import BeautifulSoup
from dateutil.parser import parse as parse_date

SUPPORTED = {
    "pdf",
    "docx",
    "txt",
    "md",
    "html",
    "srt",
    "vtt",
    "eml",
    "msg",
    "json",
    "csv",
}


def stamp(value, dayfirst=True):
    if not value:
        return None
    try:
        if re.match(r"^\d{4}-\d{2}-\d{2}", str(value)):
            from datetime import datetime

            return datetime.fromisoformat(str(value).replace("Z", "+00:00")).isoformat()
        if re.fullmatch(r"\d{10}(?:\.\d+)?", str(value)):
            from datetime import datetime, timezone

            return datetime.fromtimestamp(float(value), timezone.utc).isoformat()
        return parse_date(str(value), dayfirst=dayfirst).isoformat()
    except (ValueError, OverflowError):
        return None


def parse_source(name: str, content: bytes, dayfirst=True):
    ext = Path(name).suffix.lower().lstrip(".") or "txt"
    if ext not in SUPPORTED:
        raise ValueError("Unsupported file type: " + ext)
    segments = []
    metadata = {}

    def add(text, **meta):
        if text.strip():
            segments.append({"text": text, "index": len(segments) + 1, **meta})

    if ext == "pdf":
        import pymupdf

        with pymupdf.open(stream=content, filetype="pdf") as doc:
            if len(doc) > 1000:
                raise ValueError("PDF exceeds the 1000-page limit")
            for n, page in enumerate(doc):
                add(page.get_text(), page=n + 1)
        if not segments:
            raise ValueError(
                "This PDF has no extractable text. Local OCR is required; OCR is not installed by ChronoSync."
            )
    elif ext == "docx":
        from docx import Document

        doc = Document(io.BytesIO(content))
        for n, p in enumerate(doc.paragraphs):
            add(p.text, paragraph=n + 1)
        for n, table in enumerate(doc.tables):
            for row in table.rows:
                add(" | ".join(c.text for c in row.cells), table=n + 1)
    elif ext in {"eml", "msg"}:
        if ext == "eml":
            msg = BytesParser(policy=policy.default).parsebytes(content)
            body = msg.get_body(preferencelist=("plain", "html"))
            text = body.get_content() if body else ""
            if body and body.get_content_type() == "text/html":
                text = BeautifulSoup(text, "html.parser").get_text("\n")
            try:
                timestamp = parsedate_to_datetime(str(msg["date"])).isoformat()
            except (ValueError, TypeError):
                timestamp = None
            metadata = {
                "sender": str(msg["from"] or ""),
                "recipients": str(msg["to"] or ""),
                "subject": str(msg["subject"] or ""),
                "timestamp": timestamp,
            }
        else:
            try:
                import extract_msg
            except ImportError as exc:
                raise ValueError(
                    "MSG requires the optional extract-msg package; export as EML or install it locally."
                ) from exc
            with extract_msg.Message(content) as msg:
                text = msg.body or ""
                metadata = {
                    "sender": msg.sender,
                    "recipients": msg.to,
                    "subject": msg.subject,
                    "timestamp": stamp(msg.date, dayfirst),
                }
        add(text, **metadata, reference_kind="email timestamp")
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = (
                content.decode("utf-16")
                if content.startswith((b"\xff\xfe", b"\xfe\xff"))
                else content.decode("cp1252")
            )
        if ext in {"srt", "vtt"}:
            for block in re.split(r"\r?\n\s*\r?\n", text):
                match = re.search(
                    r"(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{3})\s*-->\s*(\d{1,2}:\d{2}(?::\d{2})?[.,]\d{3})[^\n]*\n([\s\S]+)",
                    block,
                )
                if match:
                    raw = match[3].strip()
                    voice = re.search(r"<v\s+([^>]+)>", raw)
                    speaker = (
                        voice[1]
                        if voice
                        else (
                            raw.split(":", 1)[0]
                            if re.match(r"^[A-Za-z ]{2,35}:", raw)
                            else None
                        )
                    )
                    add(
                        BeautifulSoup(raw, "html.parser").get_text(),
                        start_timestamp=match[1],
                        end_timestamp=match[2],
                        speaker=speaker,
                    )
        elif ext == "json":
            data = json.loads(text)
            rows = data if isinstance(data, list) else data.get("messages", [])
            for row in rows:
                if isinstance(row, dict):
                    add(
                        str(row.get("text", row.get("content", ""))),
                        sender=row.get("user", row.get("sender")),
                        timestamp=stamp(row.get("ts", row.get("timestamp")), dayfirst),
                        reference_kind="message timestamp",
                    )
        elif ext == "csv":
            for row in csv.DictReader(io.StringIO(text)):
                add(
                    row.get("text", row.get("message", row.get("content", ""))),
                    sender=row.get("sender"),
                    timestamp=stamp(row.get("timestamp", row.get("date")), dayfirst),
                    reference_kind="message timestamp",
                )
        elif ext == "html":
            soup = BeautifulSoup(text, "html.parser")
            for tag in soup(["script", "style"]):
                tag.decompose()
            add(soup.get_text("\n"))
        else:
            chat = re.compile(
                r"^\[?(\d{1,4}[/.-]\d{1,2}[/.-]\d{2,4}),?\s+(\d{1,2}:\d{2}(?:\s*[APap][Mm])?)\]?\s*[-–]?\s*([^:]+):\s*(.*)$"
            )
            for line in text.splitlines():
                m = chat.match(line.replace("\u200e", "").replace("\u202f", " "))
                if m:
                    add(
                        m[4],
                        sender=m[3],
                        timestamp=stamp(m[1] + " " + m[2], dayfirst),
                        reference_kind="message timestamp",
                    )
                elif (
                    segments
                    and segments[-1].get("reference_kind") == "message timestamp"
                ):
                    segments[-1]["text"] += "\n" + line
                else:
                    add(line)
    if not segments:
        raise ValueError("No readable text found in source")
    for segment in segments[:10]:
        date_header = re.search(
            r"(?im)^(?:document date|meeting date|transcript date|date):\s*(.+)$",
            segment["text"],
        )
        if date_header:
            metadata["document_date"] = stamp(date_header[1], dayfirst)
            break
    return {"segments": segments, "metadata": metadata, "type": ext}
